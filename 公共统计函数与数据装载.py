# -*- coding: utf-8 -*-
"""公共工具：统计函数与数据装载"""
import math
import numpy as np
import pandas as pd
from openpyxl import load_workbook

XLSX = r"C:\A_t\Code\数据分析\Data.xlsx"

# ---------- 分布函数（纯 Python 实现，避免额外依赖） ----------
def betacf(a, b, x):
    MAXIT, EPS, FPMIN = 300, 3e-16, 1e-300
    qab, qap, qam = a + b, a + 1.0, a - 1.0
    c = 1.0
    d = 1.0 - qab * x / qap
    if abs(d) < FPMIN: d = FPMIN
    d = 1.0 / d
    h = d
    for m in range(1, MAXIT + 1):
        m2 = 2 * m
        aa = m * (b - m) * x / ((qam + m2) * (a + m2))
        d = 1.0 + aa * d
        if abs(d) < FPMIN: d = FPMIN
        c = 1.0 + aa / c
        if abs(c) < FPMIN: c = FPMIN
        d = 1.0 / d
        h *= d * c
        aa = -(a + m) * (qab + m) * x / ((a + m2) * (qap + m2))
        d = 1.0 + aa * d
        if abs(d) < FPMIN: d = FPMIN
        c = 1.0 + aa / c
        if abs(c) < FPMIN: c = FPMIN
        d = 1.0 / d
        de = d * c
        h *= de
        if abs(de - 1.0) < EPS: break
    return h

def betai(a, b, x):
    if x <= 0: return 0.0
    if x >= 1: return 1.0
    lbeta = math.lgamma(a + b) - math.lgamma(a) - math.lgamma(b)
    bt = math.exp(lbeta + a * math.log(x) + b * math.log(1.0 - x))
    if x < (a + 1.0) / (a + b + 2.0):
        return bt * betacf(a, b, x) / a
    return 1.0 - bt * betacf(b, a, 1.0 - x) / b

def t_sf2(t, df):
    t = abs(float(t))
    return betai(df / 2.0, 0.5, df / (df + t * t))

def f_sf(f, df1, df2):
    f = float(f)
    if f <= 0: return 1.0
    return betai(df2 / 2.0, df1 / 2.0, df2 / (df2 + df1 * f))

def gammap(a, x):
    if x < 0 or a <= 0: return 0.0
    if x == 0: return 0.0
    gln = math.lgamma(a)
    if x < a + 1.0:
        ap, s, d = a, 1.0 / a, 1.0 / a
        for _ in range(500):
            ap += 1.0
            d *= x / ap
            s += d
            if abs(d) < abs(s) * 1e-15: break
        return s * math.exp(-x + a * math.log(x) - gln)
    FPMIN, b, c, d = 1e-300, x + 1.0 - a, 1e300, 1.0 / (x + 1.0 - a)
    h = d
    for i in range(1, 500):
        an = -i * (i - a)
        b += 2.0
        d = an * d + b
        if abs(d) < FPMIN: d = FPMIN
        c = b + an / c
        if abs(c) < FPMIN: c = FPMIN
        d = 1.0 / d
        de = d * c
        h *= de
        if abs(de - 1.0) < 1e-15: break
    q = math.exp(-x + a * math.log(x) - gln) * h
    return 1.0 - q

def chi2_sf(x, df):
    x = float(x)
    if x <= 0: return 1.0
    return 1.0 - gammap(df / 2.0, x / 2.0)

def cronbach_alpha(X):
    X = np.asarray(X, dtype=float)
    k = X.shape[1]
    if k < 2: return float('nan')
    item_var = X.var(axis=0, ddof=1).sum()
    total_var = X.sum(axis=1).var(ddof=1)
    if total_var == 0: return float('nan')
    return k / (k - 1.0) * (1.0 - item_var / total_var)

def alpha_if_deleted(X):
    X = np.asarray(X, dtype=float)
    return {j: cronbach_alpha(X[:, [i for i in range(X.shape[1]) if i != j]])
            for j in range(X.shape[1])}

def desc(x):
    x = np.asarray([v for v in x if v is not None and not (isinstance(v, float) and math.isnan(v))], dtype=float)
    return dict(n=len(x), mean=x.mean(), sd=x.std(ddof=1), median=float(np.median(x)),
                mn=x.min(), mx=x.max())

def corr_p(a, b):
    a = np.asarray(a, dtype=float); b = np.asarray(b, dtype=float)
    m = (~np.isnan(a)) & (~np.isnan(b))
    a, b = a[m], b[m]
    n = len(a)
    r = np.corrcoef(a, b)[0, 1]
    if abs(r) >= 1: return r, 0.0, n
    t = r * math.sqrt((n - 2) / (1 - r * r))
    return r, t_sf2(t, n - 2), n

def ols(y, Xs, names):
    y = np.asarray(y, dtype=float)
    m = ~np.isnan(y)
    for x in Xs:
        m &= ~np.isnan(np.asarray(x, dtype=float))
    yv = y[m]
    Xm = np.column_stack([np.ones(int(m.sum()))] + [np.asarray(x, dtype=float)[m] for x in Xs])
    n, p = Xm.shape
    beta, *_ = np.linalg.lstsq(Xm, yv, rcond=None)
    resid = yv - Xm @ beta
    dfres = n - p
    s2 = (resid @ resid) / dfres
    se = np.sqrt(np.diag(np.linalg.inv(Xm.T @ Xm)) * s2)
    tvals = beta / se
    pvals = [t_sf2(tv, dfres) for tv in tvals]
    sst = ((yv - yv.mean()) ** 2).sum()
    r2 = 1 - (resid @ resid) / sst
    adjr2 = 1 - (1 - r2) * (n - 1) / dfres
    F = (r2 / (p - 1)) / ((1 - r2) / dfres)
    return {'n': n, 'r2': r2, 'adjr2': adjr2, 'F': F, 'Fp': f_sf(F, p - 1, dfres), 'dfres': dfres,
            'coef': dict(zip(['常数'] + names, beta)), 'se': dict(zip(['常数'] + names, se)),
            't': dict(zip(['常数'] + names, tvals)), 'p': dict(zip(['常数'] + names, pvals))}

def anova(groups):
    groups = [np.asarray(g, dtype=float) for g in groups if len(g) > 0]
    allv = np.concatenate(groups); N = len(allv); kk = len(groups)
    gm = allv.mean()
    ssb = sum(len(g) * (g.mean() - gm) ** 2 for g in groups)
    ssw = sum(((g - g.mean()) ** 2).sum() for g in groups)
    df1, df2 = kk - 1, N - kk
    F = (ssb / df1) / (ssw / df2)
    return F, df1, df2, f_sf(F, df1, df2)

def ttest(a, b):
    a = np.asarray(a, dtype=float); b = np.asarray(b, dtype=float)
    n1, n2 = len(a), len(b)
    s1, s2 = a.var(ddof=1), b.var(ddof=1)
    sp2 = ((n1 - 1) * s1 + (n2 - 1) * s2) / (n1 + n2 - 2)
    t = (a.mean() - b.mean()) / math.sqrt(sp2 * (1 / n1 + 1 / n2))
    return t, t_sf2(t, n1 + n2 - 2), abs(a.mean() - b.mean()) / math.sqrt(sp2)

# ---------- 数据装载 ----------
def load_data():
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    data = rows[1:]
    df = pd.DataFrame(data, columns=[f"c{i}" for i in range(len(rows[0]))])
    COLS = {}
    for i in range(33):
        letter = chr(ord('A') + i) if i < 26 else 'A' + chr(ord('A') + i - 26)
        COLS[letter] = f"c{i}"
    col = lambda L: COLS[L]

    for c in [col(x) for x in ['L','M','N','O','P','Q','R','S','T','U','V','W','X','Y','Z','AA','AB','AC','AF']]:
        df[c] = pd.to_numeric(df[c], errors='coerce')

    Q7 = [col(c) for c in ['L','M','N','O','P','Q']]
    IDENT = [col(c) for c in ['R','S','T','U','V','W','X','Y','Z','AA','AB','AC']]

    comp = pd.DataFrame(index=df.index)
    comp['算法感知'] = df[[col('L'), col('N')]].mean(axis=1)
    comp['同质化感知'] = df[[col('M'), col('O'), col('P')]].mean(axis=1)
    comp['信息茧房感知'] = df[[col('N'), col('O'), col('P')]].mean(axis=1)
    comp['认知认同'] = df[[col('R'), col('S'), col('T')]].mean(axis=1)
    comp['情感认同'] = df[[col('U'), col('V'), col('W')]].mean(axis=1)
    comp['价值认同'] = df[[col('X'), col('Y'), col('Z')]].mean(axis=1)
    comp['行为认同'] = df[[col('AA'), col('AB'), col('AC')]].mean(axis=1)
    comp['主流文化认同总分'] = df[IDENT].mean(axis=1)
    comp['思政课评价'] = df[col('AF')]
    comp['数智情境感知总分'] = df[Q7].mean(axis=1)
    comp['学历阶段'] = df[col('H')]
    comp['性别'] = df[col('G')]
    comp['学科'] = df[col('I')]
    comp['学校'] = df[col('F')]
    comp['上网时长'] = df[col('J')]
    comp['AI使用'] = df[col('K')]

    ID_NAME = {col('R'):'Q8-1 清楚主流文化主要包括中华优秀传统文化、革命文化和社会主义先进文化',
               col('S'):'Q8-2 了解中华优秀传统文化的主要内容和核心精神',
               col('T'):'Q8-3 了解社会主义核心价值观的具体内容',
               col('U'):'Q8-4 为中华文化感到自豪',
               col('V'):'Q8-5 主流文化让我有归属感',
               col('W'):'Q8-6 主流文化带给我的精神力量难以替代',
               col('X'):'Q8-7 认为主流文化所倡导的价值观是正确的',
               col('Y'):'Q8-8 人生目标与主流文化倡导的价值取向一致',
               col('Z'):'Q8-9 价值选择时以主流价值观作为判断标准',
               col('AA'):'Q8-10 会主动学习主流文化相关知识',
               col('AB'):'Q8-11 会在社交平台转发或点赞主流文化内容',
               col('AC'):'Q8-12 面对歪曲历史、亵渎英雄的内容会抵制或举报'}
    Q7_NAME = {col('L'):'Q7-1 能明显感觉平台按浏览记录推送',
               col('M'):'Q7-2 刷到的内容与我的兴趣和已有观点一致',
               col('N'):'Q7-3 了解推荐算法的推送逻辑',
               col('O'):'Q7-4 很少在网上看到不同观点',
               col('P'):'Q7-5 觉得自己接收的信息单一、同质化',
               col('Q'):'Q7-6 会有意识关注不同领域账号、主动打破同质化'}
    return dict(df=df, comp=comp, col=col, COLS=COLS, Q7=Q7, IDENT=IDENT,
                ID_NAME=ID_NAME, Q7_NAME=Q7_NAME,
                SCHOOL=col('F'), SEX=col('G'), GRADE=col('H'), DISC=col('I'),
                HOURS=col('J'), AI=col('K'), FORM_MULTI=col('AD'),
                FORM_SINGLE=col('AE'), SIKE=col('AF'), POLICY=col('AG'))
